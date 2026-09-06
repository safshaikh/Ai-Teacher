import os
import re
import uuid
from typing import List, Dict, Any, Optional
from app.models import Citation

class RAGEngine:
    def __init__(self):
        self.documents: Dict[str, Dict[str, Any]] = {}
        self.chunks_db: Dict[str, List[Dict[str, Any]]] = {}

    def extract_text_from_pdf(self, file_path: str) -> List[Dict[str, Any]]:
        pages = []
        try:
            import fitz  # PyMuPDF
            doc = fitz.open(file_path)
            for page_num in range(len(doc)):
                page = doc.load_page(page_num)
                text = page.get_text("text")
                if text.strip():
                    pages.append({"page": page_num + 1, "text": text.strip()})
            doc.close()
        except Exception as e:
            # Fallback text reading if fitz isn't installed or fails
            try:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    content = f.read()
                    pages.append({"page": 1, "text": content})
            except Exception as ex:
                print(f"Error reading PDF {file_path}: {e}, {ex}")
        return pages

    def extract_text_from_docx(self, file_path: str) -> List[Dict[str, Any]]:
        sections = []
        try:
            import docx
            doc = docx.Document(file_path)
            current_section = "Introduction"
            buffer = []
            
            for para in doc.paragraphs:
                text = para.text.strip()
                if not text:
                    continue
                if para.style.name.startswith("Heading"):
                    if buffer:
                        sections.append({"page": current_section, "text": "\n".join(buffer)})
                        buffer = []
                    current_section = text
                else:
                    buffer.append(text)
            if buffer:
                sections.append({"page": current_section, "text": "\n".join(buffer)})
        except Exception as e:
            print(f"Error reading DOCX {file_path}: {e}")
        return sections

    def extract_text_from_pptx(self, file_path: str) -> List[Dict[str, Any]]:
        slides = []
        try:
            from pptx import Presentation
            prs = Presentation(file_path)
            for i, slide in enumerate(prs.slides):
                slide_text = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_text.append(shape.text.strip())
                if slide_text:
                    slides.append({"page": f"Slide {i+1}", "text": "\n".join(slide_text)})
        except Exception as e:
            print(f"Error reading PPTX {file_path}: {e}")
        return slides

    def chunk_text(self, pages_or_sections: List[Dict[str, Any]], chunk_size: int = 400, overlap: int = 50) -> List[Dict[str, Any]]:
        chunks = []
        for item in pages_or_sections:
            location = item.get("page", 1)
            text = item.get("text", "")
            
            # Split text by paragraphs/sentences
            paragraphs = re.split(r'\n\s*\n|\n', text)
            current_chunk = []
            current_len = 0
            
            for para in paragraphs:
                para = para.strip()
                if not para:
                    continue
                words = para.split()
                if current_len + len(words) > chunk_size:
                    if current_chunk:
                        chunk_str = " ".join(current_chunk)
                        chunks.append({
                            "location": str(location),
                            "text": chunk_str
                        })
                    current_chunk = words
                    current_len = len(words)
                else:
                    current_chunk.extend(words)
                    current_len += len(words)
                    
            if current_chunk:
                chunks.append({
                    "location": str(location),
                    "text": " ".join(current_chunk)
                })
        return chunks

    def process_and_index_document(self, file_path: str, original_filename: str) -> str:
        doc_id = str(uuid.uuid4())[:8]
        ext = os.path.splitext(original_filename)[1].lower()
        
        if ext == ".pdf":
            raw_data = self.extract_text_from_pdf(file_path)
        elif ext in [".docx", ".doc"]:
            raw_data = self.extract_text_from_docx(file_path)
        elif ext in [".pptx", ".ppt"]:
            raw_data = self.extract_text_from_pptx(file_path)
        else:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                raw_data = [{"page": "Main", "text": f.read()}]

        chunks = self.chunk_text(raw_data)
        
        self.documents[doc_id] = {
            "id": doc_id,
            "filename": original_filename,
            "total_chunks": len(chunks),
            "raw_sections": len(raw_data)
        }
        self.chunks_db[doc_id] = chunks
        return doc_id

    def query_citations(self, doc_id: str, query: str, top_k: int = 3) -> List[Citation]:
        if not doc_id or doc_id not in self.chunks_db:
            return []

        chunks = self.chunks_db[doc_id]
        filename = self.documents[doc_id]["filename"]
        
        query_words = set(re.findall(r'\w+', query.lower()))
        if not query_words:
            return []

        scored_chunks = []
        for chunk in chunks:
            text = chunk["text"]
            chunk_words = set(re.findall(r'\w+', text.lower()))
            overlap = len(query_words.intersection(chunk_words))
            score = overlap / max(len(query_words), 1)
            scored_chunks.append((score, chunk))

        scored_chunks.sort(key=lambda x: x[0], reverse=True)
        results = []
        
        for score, chunk in scored_chunks[:top_k]:
            if score > 0.05:
                snippet = chunk["text"][:250] + "..." if len(chunk["text"]) > 250 else chunk["text"]
                results.append(Citation(
                    source_doc=filename,
                    page_or_section=f"Page/Section {chunk['location']}",
                    snippet=snippet,
                    relevance_score=round(score, 2)
                ))

        return results

rag_engine = RAGEngine()
